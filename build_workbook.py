#!/usr/bin/env python3
"""
Build Fiix_Permissions_Reference.xlsx from the data/ files.

The data files were copied from the Fiix Help Center (permission library) on
2026-10-02. Re-run this script after editing anything in data/:

    python build_workbook.py
"""
from __future__ import annotations

import re
from collections import defaultdict
from copy import copy
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import sys

HERE = Path(__file__).parent
DATA = HERE / "data"
# Optional first argument overrides the output path.
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "Fiix_Permissions_Reference.xlsx"
RETRIEVED = "2026-10-02"
HELP = "https://helpdesk.fiixsoftware.com/hc/en-us/"

GROUPS = ["Guests", "Logistics", "Manager", "Purchasing", "Technician", "Admin"]
API_GROUPS = ["Administrators", "Managers", "Technicians", "Purchasing", "Logistics"]
CHECK = "✓"

# ---------------------------------------------------------------- styles ----
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
YES_FILL = PatternFill("solid", fgColor="C6EFCE")
YES_FONT = Font(name="Arial", size=10, bold=True, color="006100")
PART_FILL = PatternFill("solid", fgColor="FFEB9C")
NONE_FILL = PatternFill("solid", fgColor="F2F2F2")
NONE_FONT = Font(name="Arial", size=10, color="A6A6A6")
BAND_FILL = PatternFill("solid", fgColor="DDEBF7")
TITLE_FONT = Font(name="Arial", size=14, bold=True, color="1F3864")
BOLD = Font(name="Arial", size=10, bold=True)
INPUT_BLUE = Font(name="Arial", size=10, color="0000FF")
THIN = Side(style="thin", color="D9D9D9")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(wrap_text=True, vertical="top")


def header(ws, row, labels, widths=None):
    for i, text in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=text)
        c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30


def put_flag(cell, on: bool):
    cell.border = BOX
    cell.alignment = CENTER
    if on:
        cell.value, cell.fill, cell.font = CHECK, YES_FILL, YES_FONT
    else:
        cell.value, cell.font = None, NONE_FONT


# ------------------------------------------------------------- data load ----
def read_lines(name: str) -> list[str]:
    return [ln.rstrip("\n") for ln in (DATA / name).read_text(encoding="utf-8").splitlines() if ln.strip()]


def load_page_perms():
    rows, page = [], None
    for ln in read_lines("page_perms.txt"):
        if ln.startswith("#"):
            page = ln[1:]
        else:
            perm, bits = ln.rsplit("\t", 1)
            assert len(bits) == 6, ln
            rows.append((page, perm, [b == "1" for b in bits]))
    return rows


def load_menu_perms():
    out = []
    for ln in read_lines("menu_perms.tsv"):
        perm, bits = ln.split("\t")
        level = "Top menu" if "Top Menu" in perm else "Menu item"
        out.append((perm, level, [b == "1" for b in bits]))
    return out


def load_purchasing():
    out = []
    for ln in read_lines("purchasing.tsv"):
        status, perm, bits = ln.split("\t")
        out.append((status, perm, [b == "1" for b in bits]))
    return out


def load_api():
    sections, cur = {}, None
    for ln in read_lines("api_perms.txt"):
        if ln.startswith("#"):
            cur = ln[1:]
            sections[cur] = {}
        else:
            cls, bits = ln.split("\t")
            sections[cur][cls] = bits
    return sections


def split_fields(line: str, n: int) -> list[str]:
    parts = [p.strip() for p in re.split(r"\s*¦\s*", line)]
    return (parts + [""] * n)[:n]


def load_v6(name: str):
    src_name = src_url = None
    rows = []
    for ln in read_lines(name):
        if ln.startswith("@Source"):
            _, src_name, src_url = ln.split("\t")
        else:
            rows.append(split_fields(ln, 3))
    return src_name, src_url, rows


def load_feature(name: str):
    src_name = src_url = None
    mode, menu, page = None, [], []
    for ln in read_lines(name):
        if ln.startswith("@Source"):
            _, src_name, src_url = ln.split("\t")
        elif ln.startswith("@Menu"):
            mode = "menu"
        elif ln.startswith("@Page"):
            mode = "page"
        elif mode == "menu":
            menu.append(split_fields(ln, 2))
        else:
            page.append(split_fields(ln, 3))
    return src_name, src_url, menu, page


PREREQ_RE = re.compile(r"(.+?)\s+(page permissions|menu permission)\s*")


def parse_prereqs(text: str) -> list[str]:
    out = []
    for name, kind in PREREQ_RE.findall(text):
        out.append(f"{name.strip()} ({'page' if kind.startswith('page') else 'menu'})")
    return out


# ------------------------------------------------------------ dependency ----
class Edge:
    def __init__(self, area, dependent, requires, how, source, url):
        self.area, self.dependent, self.requires = area, dependent, requires
        self.how, self.source, self.url = how, source, url


def build_edges():
    edges: list[Edge] = []
    # NOTE: the v6 system role and v6 work order prerequisites are deliberately NOT included here.
    # This tenant runs v5, where those features do not exist. They live on the "v6 Reference" sheet.
    # v5 user management: what each feature needs, from Fiix's v5/v6 comparison article.
    name, url, rows = load_v6("v5_user_mgmt.txt")
    for line in read_lines("v5_user_mgmt.txt"):
        if line.startswith("@"):
            continue
        parts = [p.strip() for p in re.split(r"\s*¦\s*", line)]
        for req in parts[1:]:
            edges.append(Edge("User management (v5)", parts[0], req, "Must also be granted", name, url))
    # v5 feature articles: documented prerequisites
    area_for = {"feature_asset.txt": "Assets (v5)", "feature_sm.txt": "Scheduled maintenance (v5)",
                "feature_usermgmt.txt": "User management (v5)"}
    for fname, area in area_for.items():
        name, url, _menu, pages = load_feature(fname)
        for page, _subs, details in pages:
            if "Prerequisite permissions:" in details:
                tail = details.split("Prerequisite permissions:", 1)[1]
                for req in parse_prereqs(tail):
                    edges.append(Edge(area, f"{page} (page)", req, "Must also be granted", name, url))
    return edges


GENERAL_RULES = [
    ("Any page permission", "Matching menu permission",
     "Page permissions only work if the menu permission that leads to the page is also granted.",
     "About permissions", HELP + "articles/360049317472-About-permissions"),
    ("Any menu permission", "Matching page permissions",
     "A menu item with no page permissions shows but does nothing when clicked.",
     "About permissions", HELP + "articles/360049317472-About-permissions"),
    ("Purchasing permissions", "Matching menu and page permissions",
     "Purchasing workflow permissions only work with the related menu and page permissions.",
     "About permissions", HELP + "articles/360049317472-About-permissions"),
    ("Settings > Users and User groups screens", "Can See 'Settings' top menu, 'Users' menu item, 'User groups' menu item",
     "These menu permissions are described as mandatory for basic user management actions.",
     "User management permissions (v5)", HELP + "articles/41281346164244-User-management-permissions"),
    ("Mobile app / API access", "API permissions (Add / Find / Change / Remove per class)",
     "What a group can do in the mobile app and in apps connected through the API is set on the API Permissions tab, separately from page permissions.",
     "About permissions", HELP + "articles/360049317472-About-permissions"),
]


def build_tree(edges: list[Edge]):
    """Return ordered rows (area, level, node, also_requires, how)."""
    reqs = defaultdict(list)
    how = {}
    area = {}
    for e in edges:
        reqs[e.dependent].append(e.requires)
        how[e.dependent] = e.how
        area.setdefault(e.dependent, e.area)
        area.setdefault(e.requires, e.area)
    nodes = set(reqs) | {r for rs in reqs.values() for r in rs}
    depth_cache: dict[str, int] = {}

    def depth(n):
        if n not in depth_cache:
            depth_cache[n] = 0 if not reqs.get(n) else 1 + max(depth(r) for r in reqs[n])
        return depth_cache[n]

    popularity = defaultdict(int)
    for e in edges:
        popularity[e.requires] += 1
    parent = {}
    for n in nodes:
        if reqs.get(n):
            best = sorted(reqs[n], key=lambda r: (-depth(r), -popularity[r], r))[0]
            parent[n] = best
    children = defaultdict(list)
    for n, p in parent.items():
        children[p].append(n)
    roots = sorted((n for n in nodes if n not in parent), key=lambda n: (area[n], n))
    area_order = ["Assets (v5)", "Scheduled maintenance (v5)", "User management (v5)"]
    roots.sort(key=lambda n: (area_order.index(area[n]) if area[n] in area_order else 99, n))
    out = []

    def walk(n, level):
        also = [r for r in reqs.get(n, []) if r != parent.get(n)]
        out.append((area[n], level, n, "; ".join(also), how.get(n, "")))
        for c in sorted(children[n]):
            walk(c, level + 1)

    for r in roots:
        if children[r]:  # a root nobody hangs under is listed in the "Also requires" column instead
            walk(r, 0)
    return out


# ---------------------------------------------------------------- sheets ----
def sheet_readme(wb, sources):
    ws = wb.active
    ws.title = "Read Me"
    ws["A1"] = "Fiix permissions reference"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Compiled from the Fiix Help Center permission library on {RETRIEVED}. Default settings only."
    ws["A2"].font = Font(name="Arial", size=10, italic=True)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 110
    rows = [
        ("What this is", "A searchable copy of the default Fiix user group permissions plus every documented dependency between permissions. "
                         "Use it to see what each default group can do and what has to be switched on together. "
                         "Scoped to Fiix v5. Content that only exists in v6 is on the grey 'v6 Reference (not v5)' tab and is excluded everywhere else."),
        ("How v5 works", "Permissions live only on user groups. Go to Settings > User Groups, open a group, and use its Menu, Page, Purchasing and API "
                         "permission tabs. There are no system roles in v5, because user groups do that job. The fifth layer, role permissions, is v6 only."),
        ("What it is not", "It does not show YOUR tenant's settings. Any group you have edited or created will differ. "
                           "Fiix's own API can export your actual groups, which is a good next check."),
        ("Default groups", "Guests, Logistics, Manager, Purchasing, Technician, Admin. Source tables abbreviate these as Guest, Log., Mgr., Purch., Tech., Admin."),
        ("Sheets", ""),
        ("  Summary", "Counts per group, calculated live from the sheets below with formulas."),
        ("  Page Permissions", "Every page and sub-permission with the default groups that get it. Filter by group or page."),
        ("  Menu Permissions", "Left navigation items per default group."),
        ("  Purchasing Permissions", "Purchase order workflow steps per default group."),
        ("  API Permissions", "Per Fiix class: A=Add, F=Find, C=Change, R=Remove. These class names are the same ones the Fiix API accepts as className."),
        ("  Dependencies", "Every documented prerequisite, one row per link, plus general rules about how the permission types fit together."),
        ("  Dependency Tree", "The same links drawn as an indented tree. Each permission sits under its deepest prerequisite; any others are in 'Also requires'."),
        ("  Feature Permissions (v5)", "Descriptions of asset, scheduled maintenance, v5 work order and user management permissions."),
        ("  v6 Reference (not v5)", "System role and v6 work order permissions. Does not apply to a v5 tenant."),
        ("Cautions", ""),
        ("  Version is per feature", "Fiix rolls v6 out feature by feature, so Work Orders and User Management can each be v5 or v6. "
                                     "Fiix's two 'Which version am I using' articles show how to tell: Maintenance > Work Orders, and Settings > Users."),
        ("  Help Center how-to is v6", "Fiix's step-by-step user group articles, including 'Basic setup: Permissions, users groups, and roles', are written for v6. "
                                       "For v5, use 'About permissions' and this workbook."),
        ("  Source quirk: seats handler", "The user management permissions article calls 'User Increase Tenant Seats - UI handler' defunct. "
                                          "Fiix's v5/v6 comparison still lists it as a requirement for changing user status. Both are shown."),
        ("  API in v5", "Fiix says v5 tenants can reach all tables in the API reference, and that API access is Enterprise only."),
        ("  Source quirk: Technician menu", "The default menu table shows Technicians WITHOUT the Work Orders and Assigned Work Orders menu items, "
                                            "while the default user groups article says Technicians can see work orders. Verify in your tenant."),
        ("  Source quirk: blank page name", "One row in the page permissions table, 'Has Transfer stock button', has a blank page name in the source. "
                                            "It sits directly after Stock, so it is listed there and marked."),
        ("  Dependencies are partial", "Fiix documents prerequisites for system roles, v6 work orders, assets, scheduled maintenance and user status. "
                                       "Other pages have none published, which does not prove none exist."),
        ("  Wording", "Descriptions are lightly condensed from the Help Center text. Check the linked article for the authoritative wording."),
        ("Derived columns", "'Page type' on Page Permissions and 'Level' on Menu Permissions were added here, not taken from Fiix. "
                            "List page = has a New button. Detail page = has a Save button."),
    ]
    r = 4
    for label, text in rows:
        a = ws.cell(row=r, column=1, value=label)
        b = ws.cell(row=r, column=2, value=text)
        a.font = BOLD if not label.startswith("  ") else Font(name="Arial", size=10)
        b.alignment = WRAP
        a.alignment = Alignment(vertical="top")
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Sources").font = BOLD
    r += 1
    header(ws, r, ["Article", "URL"])
    for name, url in sources:
        r += 1
        ws.cell(row=r, column=1, value=name).alignment = WRAP
        c = ws.cell(row=r, column=2, value=url)
        c.hyperlink = url
        c.font = Font(name="Arial", size=10, color="0563C1", underline="single")
    ws.sheet_view.showGridLines = False


def sheet_page(wb, rows):
    ws = wb.create_sheet("Page Permissions")
    header(ws, 1, ["Page", "Permission"] + GROUPS + ["Groups granted", "Page type"],
           [38, 36] + [11] * 6 + [11, 13])
    by_page = defaultdict(list)
    for page, perm, _ in rows:
        by_page[page].append(perm)

    def kind(page):
        perms = by_page[page]
        if "Has 'New' Button" in perms:
            return "List page"
        if "Has 'Save' Button" in perms:
            return "Detail page"
        return "Other"

    for i, (page, perm, bits) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=page).border = BOX
        ws.cell(row=i, column=2, value=perm).border = BOX
        for j, on in enumerate(bits):
            put_flag(ws.cell(row=i, column=3 + j), on)
        c = ws.cell(row=i, column=9, value=f'=COUNTIF(C{i}:H{i},"{CHECK}")')
        c.border, c.alignment = BOX, CENTER
        k = ws.cell(row=i, column=10, value=kind(page))
        k.border, k.alignment = BOX, CENTER
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:J{len(rows) + 1}"
    return len(rows) + 1


def sheet_menu(wb, rows):
    ws = wb.create_sheet("Menu Permissions")
    header(ws, 1, ["Permission", "Level"] + GROUPS + ["Groups granted"], [46, 12] + [11] * 6 + [11])
    for i, (perm, level, bits) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=perm).border = BOX
        ws.cell(row=i, column=2, value=level).border = BOX
        for j, on in enumerate(bits):
            put_flag(ws.cell(row=i, column=3 + j), on)
        c = ws.cell(row=i, column=9, value=f'=COUNTIF(C{i}:H{i},"{CHECK}")')
        c.border, c.alignment = BOX, CENTER
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:I{len(rows) + 1}"
    return len(rows) + 1


def sheet_purchasing(wb, rows):
    ws = wb.create_sheet("Purchasing Permissions")
    header(ws, 1, ["PO status", "Permission"] + GROUPS + ["Groups granted"], [24, 34] + [11] * 6 + [11])
    for i, (status, perm, bits) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=status).border = BOX
        ws.cell(row=i, column=2, value=perm).border = BOX
        for j, on in enumerate(bits):
            put_flag(ws.cell(row=i, column=3 + j), on)
        c = ws.cell(row=i, column=9, value=f'=COUNTIF(C{i}:H{i},"{CHECK}")')
        c.border, c.alignment = BOX, CENTER
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:I{len(rows) + 1}"
    return len(rows) + 1


def letters(bits: str) -> str:
    s = "".join(ch for ch, b in zip("AFCR", bits) if b == "1")
    return s or "-"


def sheet_api(wb, sections):
    ws = wb.create_sheet("API Permissions")
    header(ws, 1, ["Fiix class"] + API_GROUPS + ["Groups with any access"], [30] + [15] * 5 + [14])
    classes = list(sections["Administrators"].keys())
    for g in API_GROUPS:
        assert list(sections[g].keys()) == classes, g
    for i, cls in enumerate(classes, start=2):
        ws.cell(row=i, column=1, value=cls).border = BOX
        for j, g in enumerate(API_GROUPS):
            txt = letters(sections[g][cls])
            c = ws.cell(row=i, column=2 + j, value=txt)
            c.border, c.alignment = BOX, CENTER
            if txt == "AFCR":
                c.fill, c.font = YES_FILL, YES_FONT
            elif txt == "-":
                c.fill, c.font = NONE_FILL, NONE_FONT
            else:
                c.fill = PART_FILL
        c = ws.cell(row=i, column=7, value=f'=COUNTIF(B{i}:F{i},"<>-")')
        c.border, c.alignment = BOX, CENTER
    n = len(classes) + 1
    note = n + 2
    ws.cell(row=note, column=1, value="Key").font = BOLD
    ws.cell(row=note + 1, column=1, value="A = Add (create)   F = Find (read)   C = Change (edit)   R = Remove (delete)   - = no access")
    ws.cell(row=note + 2, column=1, value="Guests have no API permissions. Green = full access, amber = partial, grey = none.")
    ws.cell(row=note + 3, column=1, value="Class names are the same values you pass as className in the Fiix API (see fiix-api/fiix_query.py).")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:G{n}"
    return n


def sheet_dependencies(wb, edges):
    ws = wb.create_sheet("Dependencies")
    header(ws, 1, ["Area", "Permission", "Requires", "How enforced", "Source article"], [26, 44, 44, 24, 40])
    r = 2
    for e in edges:
        ws.cell(row=r, column=1, value=e.area)
        ws.cell(row=r, column=2, value=e.dependent)
        ws.cell(row=r, column=3, value=e.requires)
        ws.cell(row=r, column=4, value=e.how)
        s = ws.cell(row=r, column=5, value=e.source)
        s.hyperlink = e.url
        s.font = Font(name="Arial", size=10, color="0563C1", underline="single")
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = BOX
            ws.cell(row=r, column=c).alignment = WRAP
        r += 1
    for dep, req, why, src, url in GENERAL_RULES:
        ws.cell(row=r, column=1, value="General rule")
        ws.cell(row=r, column=2, value=dep)
        ws.cell(row=r, column=3, value=req)
        ws.cell(row=r, column=4, value="Both required")
        s = ws.cell(row=r, column=5, value=src)
        s.hyperlink = url
        s.font = Font(name="Arial", size=10, color="0563C1", underline="single")
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = BOX
            ws.cell(row=r, column=c).alignment = WRAP
            if c < 5:
                ws.cell(row=r, column=c).fill = BAND_FILL
        ws.cell(row=r, column=4).comment = None
        r += 1
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{r - 1}"
    return r - 1


def sheet_tree(wb, edges):
    ws = wb.create_sheet("Dependency Tree")
    header(ws, 1, ["Area", "Permission (indented = depends on the one above)", "Also requires", "How enforced"],
           [26, 70, 54, 24])
    r = 2
    for area, level, node, also, how in build_tree(edges):
        a = ws.cell(row=r, column=1, value=area)
        if level == 0:
            a.font = BOLD
        prefix = "" if level == 0 else "└─ "
        p = ws.cell(row=r, column=2, value=prefix + node)
        p.alignment = Alignment(indent=level * 2, vertical="top")
        if level == 0:
            p.font = BOLD
        ws.cell(row=r, column=3, value=also or None).alignment = WRAP
        ws.cell(row=r, column=4, value=how or None)
        for c in range(1, 5):
            ws.cell(row=r, column=c).border = BOX
        r += 1
    ws.freeze_panes = "A2"
    return r - 1


def sheet_v6(wb):
    ws = wb.create_sheet("v6 Reference (not v5)")
    ws.sheet_properties.tabColor = "A6A6A6"
    ws.merge_cells("A1:E1")
    ban = ws["A1"]
    ban.value = ("NOT APPLICABLE TO A v5 TENANT. System roles and v6 work order permissions only exist in Fiix v6. "
                 "Kept for reference in case you move to v6.")
    ban.font = Font(name="Arial", size=10, bold=True, color="9C0006")
    ban.fill = PatternFill("solid", fgColor="FFC7CE")
    ban.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    header(ws, 2, ["Area", "Permission", "What it allows", "Notes and prerequisites", "Source article"], [20, 40, 70, 60, 36])
    r = 3
    for fname, area in (("sysroles.txt", "System roles (v6)"), ("wo_v6.txt", "Work orders (v6)")):
        name, url, rows = load_v6(fname)
        for perm, beh, extra in rows:
            vals = [area, perm, beh, extra]
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=c, value=v or None)
                cell.alignment, cell.border = WRAP, BOX
            s = ws.cell(row=r, column=5, value=name)
            s.hyperlink = url
            s.font = Font(name="Arial", size=10, color="0563C1", underline="single")
            s.border, s.alignment = BOX, WRAP
            r += 1
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:E{r - 1}"
    return r - 1


def sheet_features(wb):
    ws = wb.create_sheet("Feature Permissions (v5)")
    header(ws, 1, ["Area", "Type", "Page / permission", "Sub-permissions", "What it allows", "Prerequisites", "Source article"],
           [24, 8, 36, 40, 80, 46, 34])
    r = 2
    for fname, area in (("feature_wo_v5.txt", "Work orders (v5)"), ("feature_asset.txt", "Assets (v5)"),
                        ("feature_sm.txt", "Scheduled maintenance (v5)"), ("feature_usermgmt.txt", "User management (v5)")):
        name, url, menu, pages = load_feature(fname)
        entries = [("Menu", m[0], "", m[1]) for m in menu] + [("Page", p[0], p[1], p[2]) for p in pages]
        for typ, perm, subs, details in entries:
            prereq = ""
            if "Prerequisite permissions:" in details:
                details, tail = details.split("Prerequisite permissions:", 1)
                prereq = "; ".join(parse_prereqs(tail))
            vals = [area, typ, perm, subs or None, details.strip() or None, prereq or None]
            for c, v in enumerate(vals, start=1):
                cell = ws.cell(row=r, column=c, value=v)
                cell.alignment, cell.border = WRAP, BOX
            s = ws.cell(row=r, column=7, value=name)
            s.hyperlink = url
            s.font = Font(name="Arial", size=10, color="0563C1", underline="single")
            s.border, s.alignment = BOX, WRAP
            r += 1
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:G{r - 1}"
    return r - 1


def sheet_summary(wb, page_last, menu_last, purch_last, api_last):
    ws = wb.create_sheet("Summary", 1)
    ws["A1"] = "Default permissions per group"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "All numbers are formulas over the other sheets. They update if you edit those sheets."
    ws["A2"].font = Font(name="Arial", size=10, italic=True)
    header(ws, 4, ["Group", "Menu permissions", "Page permissions", "Purchasing permissions", "Total granted", "Share of all page permissions"],
           [22, 18, 18, 22, 16, 22])
    for i, g in enumerate(GROUPS):
        col = get_column_letter(3 + i)
        r = 5 + i
        ws.cell(row=r, column=1, value=g).font = BOLD
        ws.cell(row=r, column=2, value=f"=COUNTIF('Menu Permissions'!{col}2:{col}{menu_last},\"{CHECK}\")")
        ws.cell(row=r, column=3, value=f"=COUNTIF('Page Permissions'!{col}2:{col}{page_last},\"{CHECK}\")")
        ws.cell(row=r, column=4, value=f"=COUNTIF('Purchasing Permissions'!{col}2:{col}{purch_last},\"{CHECK}\")")
        ws.cell(row=r, column=5, value=f"=SUM(B{r}:D{r})")
        ws.cell(row=r, column=6, value=f"=C{r}/$C$11").number_format = "0.0%"
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = BOX
            if c > 1:
                ws.cell(row=r, column=c).alignment = CENTER
    r = 11
    ws.cell(row=r, column=1, value="Total available").font = BOLD
    ws.cell(row=r, column=2, value=f"=COUNTA('Menu Permissions'!A2:A{menu_last})")
    ws.cell(row=r, column=3, value=f"=COUNTA('Page Permissions'!B2:B{page_last})")
    ws.cell(row=r, column=4, value=f"=COUNTA('Purchasing Permissions'!B2:B{purch_last})")
    ws.cell(row=r, column=5, value=f"=SUM(B{r}:D{r})")
    for c in range(1, 6):
        ws.cell(row=r, column=c).border = BOX
        ws.cell(row=r, column=c).fill = BAND_FILL
        ws.cell(row=r, column=c).font = BOLD
        if c > 1:
            ws.cell(row=r, column=c).alignment = CENTER

    ws["A15"] = "API access per group (Fiix classes)"
    ws["A15"].font = Font(name="Arial", size=12, bold=True, color="1F3864")
    header(ws, 16, ["Group", "Classes with any access", "Classes with Find", "Classes with any write (Add, Change or Remove)",
                    "Classes with full access", "Total classes"])
    ws.row_dimensions[16].height = 44
    for i, g in enumerate(API_GROUPS):
        col = get_column_letter(2 + i)
        rng = f"'API Permissions'!{col}2:{col}{api_last}"
        r = 17 + i
        ws.cell(row=r, column=1, value=g).font = BOLD
        ws.cell(row=r, column=2, value=f'=COUNTIF({rng},"<>-")')
        ws.cell(row=r, column=3, value=f'=COUNTIF({rng},"*F*")')
        ws.cell(row=r, column=4, value=(f'=SUMPRODUCT(--((ISNUMBER(SEARCH("A",{rng}))+ISNUMBER(SEARCH("C",{rng}))'
                                        f'+ISNUMBER(SEARCH("R",{rng})))>0))'))
        ws.cell(row=r, column=5, value=f'=COUNTIF({rng},"AFCR")')
        ws.cell(row=r, column=6, value=f"=COUNTA('API Permissions'!A2:A{api_last})")
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = BOX
            if c > 1:
                ws.cell(row=r, column=c).alignment = CENTER
    ws.cell(row=23, column=1, value="Guests have no API permissions and are not listed.").font = Font(name="Arial", size=10, italic=True)
    ws.sheet_view.showGridLines = False


def normalize_fonts(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                f = copy(cell.font)
                if f.name != "Arial":
                    f = Font(name="Arial", size=10 if f.size in (None, 11) else f.size, bold=f.bold,
                             italic=f.italic, color=f.color, underline=f.underline)
                    cell.font = f


def main():
    sources = [
        ("About permissions", HELP + "articles/360049317472-About-permissions"),
        ("About default user groups", HELP + "articles/18165584633876-About-default-user-groups"),
        ("Which version of User Management am I using?", HELP + "articles/13012571221396-Which-version-of-User-Management-am-I-using"),
        ("Which version of Work Orders am I using?", HELP + "articles/7884592276756-Which-version-of-Work-Orders-am-I-using"),
        ("Compare user management v5 and v6 features", HELP + "articles/21769760933652-Compare-user-management-v5-and-v6-features"),
        ("Default user group permissions - page permissions", HELP + "articles/360049275232-Default-user-group-permissions-page-permissions"),
        ("Default user group permissions - menu permissions", HELP + "articles/360049707091-Default-user-group-permissions-menu-permissions"),
        ("Default user group permissions - purchasing", HELP + "articles/360049317492-Default-user-group-permissions-purchasing"),
        ("Default user group permissions - API permissions", HELP + "articles/360049795251-Default-user-group-permissions-API-permissions"),
        ("Work orders permissions (v5)", HELP + "articles/40598772915988-Work-orders-permissions"),
        ("Asset permissions", HELP + "articles/42176625317652-Asset-permissions"),
        ("Scheduled maintenance permissions", HELP + "articles/42451817320340-Scheduled-maintenance-permissions"),
        ("User management permissions", HELP + "articles/41281346164244-User-management-permissions"),
        ("Permission library (section index)", HELP + "sections/360010366012-Permission-library"),
        ("[v6 only] System role permissions", HELP + "articles/25255386559252-System-role-permissions"),
        ("[v6 only] Work orders permissions (v6)", HELP + "articles/27945783675924-Work-orders-permissions-v6"),
    ]
    wb = Workbook()
    sheet_readme(wb, sources)
    page_last = sheet_page(wb, load_page_perms())
    menu_last = sheet_menu(wb, load_menu_perms())
    purch_last = sheet_purchasing(wb, load_purchasing())
    api_last = sheet_api(wb, load_api())
    edges = build_edges()
    sheet_dependencies(wb, edges)
    sheet_tree(wb, edges)
    sheet_features(wb)
    sheet_v6(wb)
    sheet_summary(wb, page_last, menu_last, purch_last, api_last)
    normalize_fonts(wb)
    wb.save(OUT)
    print(f"Saved {OUT}  |  page rows {page_last - 1}, menu {menu_last - 1}, purchasing {purch_last - 1}, "
          f"api classes {api_last - 1}, dependency edges {len(edges)}")


if __name__ == "__main__":
    main()
